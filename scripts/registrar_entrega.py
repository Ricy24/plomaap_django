#!/usr/bin/env python
# -*- coding: utf-8 -*-
"""Registra en el Excel la entrega de un integrante, solo si cumple el contrato.

Uso:
    python scripts/registrar_entrega.py --integrante sebastian
    python scripts/registrar_entrega.py --integrante kevin

El script:
  1. valida el archivo contra docs/pruebas/CONTRATO_PRUEBAS_PENDIENTES.md;
  2. ejecuta las pruebas del archivo;
  3. solo si todo pasa, agrega las 10 filas al Excel de entrega.

Si algo falla, el Excel no se modifica.
"""
from __future__ import annotations

import argparse
import ast
import re
import subprocess
import sys
from copy import copy
from pathlib import Path

from openpyxl import load_workbook
from openpyxl.formatting.formatting import ConditionalFormattingList
from openpyxl.formatting.rule import CellIsRule
from openpyxl.styles import Font, PatternFill
from openpyxl.worksheet.datavalidation import DataValidation

RAIZ = Path(__file__).resolve().parent.parent
EQUIPO = RAIZ / "api" / "tests" / "unit" / "equipo"
EXCEL = RAIZ / "docs" / "pruebas" / "Plan_de_Pruebas_PlomApp_SENA_entrega_equipo.xlsx"

INTEGRANTES = {
    "sebastian": "test_integrante_2_seguridad.py",
    "kevin": "test_integrante_4_datos.py",
}

sys.path.insert(0, str(RAIZ / "scripts"))
from check_entrega_equipo import SPEC, validar  # noqa: E402


def _campo(doc: str, etiqueta: str) -> str:
    for linea in doc.splitlines():
        if linea.strip().startswith(etiqueta):
            return linea.split(etiqueta, 1)[1].strip()
    return ""


def _pruebas_del_archivo(archivo: Path) -> list[dict]:
    arbol = ast.parse(archivo.read_text(encoding="utf-8"))
    datos = []
    for nodo in arbol.body:
        if not isinstance(nodo, (ast.FunctionDef, ast.AsyncFunctionDef)):
            continue
        if not nodo.name.startswith("test_"):
            continue
        doc = ast.get_docstring(nodo) or ""
        datos.append(
            {
                "nombre": nodo.name,
                "caso": _campo(doc, "Caso:"),
                "entrada": _campo(doc, "Entrada:"),
                "esperado": _campo(doc, "Esperado:"),
            }
        )
    return datos


def _ejecutar_pruebas(archivo: Path) -> bool:
    comando = [
        sys.executable,
        "-m",
        "pytest",
        f"api/tests/unit/equipo/{archivo.name}",
        "-q",
    ]
    resultado = subprocess.run(comando, cwd=RAIZ)
    return resultado.returncode == 0


def _ultimo_numero(ws) -> int:
    for fila in range(ws.max_row, 7, -1):
        valor = ws.cell(fila, 1).value
        if isinstance(valor, str):
            coincidencia = re.match(r"UNIT-(\d+)", valor)
            if coincidencia:
                return int(coincidencia.group(1))
    return 0


def _actualizar_formulas(ws, control, total: int) -> None:
    ultima = 7 + total
    for fila in ws.iter_rows():
        for celda in fila:
            if isinstance(celda.value, str) and celda.value.startswith("="):
                celda.value = re.sub(r"A8:A\d+", f"A8:A{ultima}", celda.value)
                celda.value = re.sub(r"I8:I\d+", f"I8:I{ultima}", celda.value)
                celda.value = re.sub(r"J8:J\d+", f"J8:J{ultima}", celda.value)

    for fila in control.iter_rows():
        for celda in fila:
            if isinstance(celda.value, str) and celda.value.startswith("="):
                celda.value = re.sub(
                    r"'Pruebas Unitarias'!A8:A\d+",
                    f"'Pruebas Unitarias'!A8:A{ultima}",
                    celda.value,
                )
                celda.value = re.sub(
                    r"'Pruebas Unitarias'!I8:I\d+",
                    f"'Pruebas Unitarias'!I8:I{ultima}",
                    celda.value,
                )


def _rehacer_validaciones(ws, total: int) -> None:
    ultima = 7 + total

    ws.data_validations.dataValidation = []
    ws.conditional_formatting = ConditionalFormattingList()

    dv_estado = DataValidation(
        type="list",
        formula1='"Exitoso,Falló,Bloqueado,Pendiente"',
        allow_blank=True,
    )
    dv_estado.add(f"I8:I{ultima}")
    ws.add_data_validation(dv_estado)

    dv_sev = DataValidation(type="list", formula1='"Alta,Media,Baja"', allow_blank=True)
    dv_sev.add(f"J8:J{ultima}")
    ws.add_data_validation(dv_sev)

    ws.conditional_formatting.add(
        f"I8:I{ultima}",
        CellIsRule(
            operator="equal",
            formula=['"Exitoso"'],
            fill=PatternFill("solid", fgColor="FFC6EFCE"),
            font=Font(color="FF006100"),
        ),
    )


def registrar(clave: str) -> int:
    if clave not in INTEGRANTES:
        print(f"[registro] Integrante desconocido: {clave}")
        return 1

    archivo = EQUIPO / INTEGRANTES[clave]
    regla = SPEC[archivo.name]

    if not archivo.exists():
        print(f"[registro] No existe {archivo.name}. Nada que registrar.")
        return 1

    print(f"[registro] Validando el contrato de {archivo.name} ...")
    errores = validar(archivo.name)
    if errores:
        print("[registro] La entrega no cumple el contrato:")
        for error in errores:
            print(f"  - {error}")
        print("[registro] No se modificó el Excel.")
        return 1

    print("[registro] Ejecutando las pruebas ...")
    if not _ejecutar_pruebas(archivo):
        print("[registro] Las pruebas no pasan. No se modificó el Excel.")
        return 1

    pruebas = _pruebas_del_archivo(archivo)
    if len(pruebas) != 10:
        print(f"[registro] Se esperaban 10 pruebas y hay {len(pruebas)}.")
        return 1

    wb = load_workbook(EXCEL)
    ws = wb["Pruebas Unitarias"]
    control = wb["Control y Resumen"]

    # Si ya estaba registrado, no duplicar.
    ya = {
        str(ws.cell(fila, 11).value)
        for fila in range(8, ws.max_row + 1)
        if ws.cell(fila, 11).value
    }
    if any(archivo.name in evidencia for evidencia in ya):
        print(f"[registro] {regla['responsable']} ya tenía su entrega registrada.")
        return 0

    ultima_actual = ws.max_row
    numero = _ultimo_numero(ws)
    plantilla = ultima_actual

    for indice, prueba in enumerate(pruebas, start=1):
        fila = ultima_actual + indice
        numero += 1
        valores = {
            1: f"UNIT-{numero:03d}",
            2: prueba["caso"] or prueba["nombre"],
            3: regla["precondiciones"],
            4: prueba["entrada"],
            5: (
                f"1. Ejecutar: pytest "
                f"api\\tests\\unit\\equipo\\{archivo.name}::{prueba['nombre']}\n"
                f"2. Verificar el resultado esperado."
            ),
            6: prueba["esperado"],
            7: "Aprobado: passed en la suite api/tests/unit/equipo.",
            8: "Caja Blanca (unitaria)",
            9: "Exitoso",
            10: regla["severidad"],
            11: (
                f"api\\tests\\unit\\equipo\\{archivo.name}::"
                f"{prueba['nombre']} -> 1 passed"
            ),
            12: f"Sin defectos. Entrega del integrante: {regla['responsable']}.",
            13: regla["responsable"],
        }

        for columna, valor in valores.items():
            celda = ws.cell(fila, columna, valor)
            celda._style = copy(ws.cell(plantilla, columna)._style)

        ws.row_dimensions[fila].height = 78

    total = numero
    _actualizar_formulas(ws, control, total)
    _rehacer_validaciones(ws, total)

    ws["A6"] = (
        f"ENTREGA ACTUAL: {total} PRUEBAS UNITARIAS registradas "
        f"(UNIT-001 a UNIT-{total:03d}). Ejecución real: {total} passed. "
        "Cada fila indica el integrante responsable."
    )

    equipo = control["C6"].value or ""
    if regla["responsable"] not in equipo:
        control["C6"] = f"{equipo}, {regla['responsable']}".strip(", ")

    wb.calculation.fullCalcOnLoad = True
    wb.save(EXCEL)

    print(
        f"[registro] OK: se agregaron 10 filas para {regla['responsable']} "
        f"(hasta UNIT-{total:03d})."
    )
    return 0


def main() -> int:
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument(
        "--integrante",
        required=True,
        choices=sorted(INTEGRANTES),
        help="Integrante que entrega.",
    )
    args = parser.parse_args()
    return registrar(args.integrante)


if __name__ == "__main__":
    sys.exit(main())
